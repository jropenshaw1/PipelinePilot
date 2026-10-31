"""ADR-015 §1: master resume parsing and pool loading (synthetic fixtures only)."""
import pytest
from docx import Document

import fit_analysis_engine as fae
from fit_analysis_engine import TailoringLoadError

ROLES = [
    ("Director, Platform Engineering", "Example Corp | Phoenix, AZ | 2020 - 2024",
     ["Built a platform team from zero.", "Cut cloud spend 20% in one year."]),
    ("Manager, Systems", "Sample Inc. | Tempe, AZ | 2015 - 2020",
     ["Ran release management for two platforms."]),
]


def make_master(path, roles=ROLES, table=False, split_bullet=False):
    d = Document()
    d.add_paragraph("PAT EXAMPLE")
    c = d.add_paragraph("555-0100 | ")
    c.add_run("pat@example.com")  # locked line: several runs allowed
    d.add_paragraph("Technology Leader | Platforms")
    d.add_paragraph("Cloud | Cost | Operations")
    d.add_paragraph("PROFILE SUMMARY")
    d.add_paragraph("Leader who builds reliable platforms.")
    d.add_paragraph("Founded two practices.", style="List Paragraph")
    b = d.add_paragraph("Owned a $5M budget.", style="List Paragraph")
    if split_bullet:
        b.add_run(" Extra run.")
    d.add_paragraph("CORE COMPETENCIES")
    d.add_paragraph("Leadership: Strategy | Team Building")
    d.add_paragraph("Finance: FinOps | Budgeting")
    d.add_paragraph("PROFESSIONAL EXPERIENCE")
    for title, company, bullets in roles:
        d.add_paragraph(title)
        d.add_paragraph(company)
        for t in bullets:
            d.add_paragraph(t, style="List Paragraph")
    p = d.add_paragraph("TECHNICAL PORTFOLIO ")
    p.add_run("(Open Source)")
    d.add_paragraph("Built a tool.", style="List Paragraph")
    d.add_paragraph("EDUCATION")
    d.add_paragraph("Example University")
    if table:
        d.add_table(rows=1, cols=1)
    d.save(str(path))
    return path


def test_parse_inventory(tmp_path):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    assert inv["headline"]["text"] == "Technology Leader | Platforms"
    assert inv["subtitle"]["text"] == "Cloud | Cost | Operations"
    assert inv["summary_paragraph"]["text"].startswith("Leader")
    assert [b["id"] for b in inv["summary_bullets"]] == ["M:S:0", "M:S:1"]
    assert inv["competencies"][1] == {"index": inv["competencies"][1]["index"],
                                      "label": "Finance", "terms": ["FinOps", "Budgeting"]}
    assert [r["title"] for r in inv["roles"]] == [ROLES[0][0], ROLES[1][0]]
    assert [b["id"] for b in inv["roles"][0]["bullets"]] == ["M:0:0", "M:0:1"]
    assert [b["id"] for b in inv["roles"][1]["bullets"]] == ["M:1:0"]
    assert len(inv["sha256"]) == 64
    # locked sections are not inventoried
    assert all("tool" not in b["text"] for r in inv["roles"] for b in r["bullets"])


def test_paragraph_indexes_point_at_the_text(tmp_path):
    path = make_master(tmp_path / "m.docx")
    inv = fae.parse_master_resume(path)
    paras = Document(str(path)).paragraphs
    b = inv["roles"][0]["bullets"][1]
    assert paras[b["index"]].text == b["text"]


@pytest.mark.parametrize("kwargs, reason", [
    ({"table": True}, "tables"),
    ({"split_bullet": True}, "text runs"),
    ({"roles": ROLES + [ROLES[1]]}, "repeated role identity"),
    ({"roles": [("Director", "Example Corp | Phoenix, AZ", ["x"])]}, "no dates"),
])
def test_parse_fails_closed(tmp_path, kwargs, reason):
    with pytest.raises(TailoringLoadError, match=reason):
        fae.parse_master_resume(make_master(tmp_path / "m.docx", **kwargs))


ROLE_MAP = """\
period-a:
  title: "Director, Platform Engineering"
  company_line: "Example Corp | Phoenix, AZ | 2020 - 2024"
period-b:
  title: "Manager, Systems"
  company_line: "Sample Inc. | Tempe, AZ | 2015 - 2020"
"""


def item(item_id, period, polished):
    return (f"### {item_id}\n\n- **role_period:** {period}\n- **raw_capture:** words\n"
            f"- **polished_resume_item:** {polished}\n- **notes:** none\n\n")


def make_pool(tmp_path, files, role_map=ROLE_MAP):
    pool = tmp_path / "pool"
    pool.mkdir()
    if role_map is not None:
        (pool / "role_map.yaml").write_text(role_map, encoding="utf-8")
    (pool / "SCHEMA.md").write_text("## Pool Items\n\n" + item("schema-x", "period-a", "Ignored."),
                                    encoding="utf-8")
    (pool / "docs").mkdir()
    (pool / "docs" / "extra.md").write_text("## Pool Items\n\n" + item("sub-x", "period-a", "Ignored."),
                                            encoding="utf-8")
    for name, body in files.items():
        (pool / name).write_text("# Period\n\nIntro.\n\n## Pool Items\n\n" + body, encoding="utf-8")
    return pool


def test_pool_loads_and_binds(tmp_path):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    pool = make_pool(tmp_path, {
        "a.md": item("a-one", "period-a", "Built a thing.") + item("a-two", "period-a", "(not yet captured)"),
        "b.md": item("b-one", "period-b", "Ran a thing.") + item("x-one", "cross-career", "Led across eras."),
    })
    out = fae.load_resume_pool(pool, inv)
    got = {i["item_id"]: i["role_index"] for i in out["items"]}
    assert got == {"a-one": 0, "b-one": 1, "x-one": None}  # placeholder skipped; unmapped = summary-only
    assert out["role_map"] == {"period-a": 0, "period-b": 1}


def test_pool_items_outside_pool_items_section_are_ignored(tmp_path):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    pool = make_pool(tmp_path, {})
    (pool / "c.md").write_text(item("early", "period-a", "Before the section.") + "## Pool Items\n\n",
                               encoding="utf-8")
    assert fae.load_resume_pool(pool, inv)["items"] == []


def test_pool_duplicate_item_id_fails(tmp_path):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    pool = make_pool(tmp_path, {"a.md": item("dup", "period-a", "One."),
                                "b.md": item("dup", "period-b", "Two.")})
    with pytest.raises(TailoringLoadError, match="duplicate pool item_id 'dup'"):
        fae.load_resume_pool(pool, inv)


@pytest.mark.parametrize("role_map, reason", [
    (None, "role_map.yaml not found"),
    ("period-a:\n  title: Nobody\n  company_line: Nowhere | 2000\n", "matches 0 master roles"),
    ("period-a: just a string\n", "needs title and company_line"),
    ("- a list\n", "must map role periods"),
    ("period-a: [unclosed\n", "not valid YAML"),
])
def test_role_map_fails_closed(tmp_path, role_map, reason):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    pool = make_pool(tmp_path, {}, role_map=role_map)
    with pytest.raises(TailoringLoadError, match=reason):
        fae.load_resume_pool(pool, inv)


def test_missing_pool_folder_fails(tmp_path):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    with pytest.raises(TailoringLoadError, match="not found"):
        fae.load_resume_pool(tmp_path / "nope", inv)


def test_in_summary_flag_excludes_role(tmp_path):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    rm = ROLE_MAP + "  in_summary: false\n"  # applies to period-b (Sample Inc.)
    pool = make_pool(tmp_path, {}, role_map=rm)
    out = fae.load_resume_pool(pool, inv)
    assert out["summary_excluded_roles"] == {1}


def test_in_summary_defaults_to_true(tmp_path):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    out = fae.load_resume_pool(make_pool(tmp_path, {}), inv)
    assert out["summary_excluded_roles"] == set()


@pytest.mark.parametrize("extra, reason", [
    ("  in_summary: no-thanks\n", "must be true or false"),
    ('period-c:\n  title: "Manager, Systems"\n  company_line: "Sample Inc. | Tempe, AZ | 2015 - 2020"\n'
     "  in_summary: false\n", "disagree on in_summary"),
])
def test_in_summary_flag_fails_closed(tmp_path, extra, reason):
    inv = fae.parse_master_resume(make_master(tmp_path / "m.docx"))
    pool = make_pool(tmp_path, {}, role_map=ROLE_MAP + extra)
    with pytest.raises(TailoringLoadError, match=reason):
        fae.load_resume_pool(pool, inv)
