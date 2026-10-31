"""ADR-015 §5.1 and §7: render, trace, recoverable publication (synthetic fixtures only)."""
import json

import pytest
from docx import Document

import fit_analysis_engine as fae
from test_resume_sources import make_master

RES, TR = "Resume_Acme_Director.docx", "ResumeTrace_Acme_Director.md"


def tailored_res(inv):
    pool = {"items": [], "summary_excluded_roles": set()}
    t = {
        "headline": "Technology Leader | Cloud",
        "subtitle": inv["subtitle"]["text"],
        "summary_paragraph": inv["summary_paragraph"]["text"],
        "summary_bullets": [{"text": b["text"], "source_ids": [b["id"]]} for b in inv["summary_bullets"]][:1],
        "competencies": [list(c["terms"]) for c in inv["competencies"]],
        "roles": [{"role_index": r["role_index"],
                   "bullets": [{"text": b["text"], "source_ids": [b["id"]]} for b in r["bullets"]][::-1]}
                  for r in inv["roles"]],
    }
    res = fae.validate_tailoring(inv, pool, t)
    fae.apply_role_limits(inv, pool, res)
    fae.apply_budget(inv, res)
    return pool, res


def texts(path):
    return [p.text for p in Document(str(path)).paragraphs if p.text.strip()]


# ── §7 render ──

def test_render_changes_only_editable_text_and_stamps_run_id(tmp_path):
    master = make_master(tmp_path / "m.docx")
    inv = fae.parse_master_resume(master)
    _, res = tailored_res(inv)
    out = fae.render_tailored_resume(master, inv, res, tmp_path / "out.docx", "20261031T000000Z-0123abcd", "tailored")
    before, after = texts(master), texts(out)
    assert after[2] == "Technology Leader | Cloud"
    assert after[0:2] == before[0:2]                      # name and contact untouched
    assert "Owned a $5M budget." not in after             # summary trimmed to one bullet
    assert after.index("Cut cloud spend 20% in one year.") < after.index("Built a platform team from zero.")
    assert after[-4:] == before[-4:]                      # portfolio, education untouched
    assert fae.resume_run_id(out) == "20261031T000000Z-0123abcd"


def test_render_keeps_bullet_formatting_and_can_add_bullets(tmp_path):
    master = make_master(tmp_path / "m.docx")
    inv = fae.parse_master_resume(master)
    _, res = tailored_res(inv)
    res["summary_bullets"].append({"text": "A third summary point.", "source_ids": ["M:S:0"], "status": "retuned"})
    res["summary_bullets"].append({"text": "A fourth summary point.", "source_ids": ["M:S:1"], "status": "retuned"})
    out = fae.render_tailored_resume(master, inv, res, tmp_path / "o.docx", "20261031T000000Z-0123abcd", "partial")
    doc = Document(str(out))
    added = [p for p in doc.paragraphs if p.text in ("A third summary point.", "A fourth summary point.")]
    assert [p.text for p in added] == ["A third summary point.", "A fourth summary point."]
    assert all(p.style.name == "List Paragraph" for p in added)
    reparsed = fae.parse_master_resume(out)              # still a well-formed master shape
    assert len(reparsed["summary_bullets"]) == 3


def test_master_copy_is_unchanged_but_stamped(tmp_path):
    master = make_master(tmp_path / "m.docx")
    out = fae.master_copy_resume(master, tmp_path / "c.docx", "20261031T000000Z-0123abcd")
    assert texts(out) == texts(master)
    assert fae.resume_run_id(out) == "20261031T000000Z-0123abcd"


# ── trace ──

def test_trace_lists_retuned_items_first_with_sources(tmp_path):
    master = make_master(tmp_path / "m.docx")
    inv = fae.parse_master_resume(master)
    pool, res = tailored_res(inv)
    tr = fae.write_resume_trace(tmp_path / TR, run_id="20261031T000000Z-0123abcd", state="tailored",
                                company="Acme", role="Director", model="claude-test-1",
                                inventory=inv, pool=pool, res=res).read_text(encoding="utf-8")
    assert fae.trace_run_id(tmp_path / TR) == "20261031T000000Z-0123abcd"
    assert tr.index("## Retuned items") < tr.index("## Changes the checks made") < tr.index("## Published resume")
    assert "### headline" in tr and "- master: Technology Leader | Platforms" in tr
    assert "**retuned items:** 1" in tr


def test_master_copy_trace_names_the_reason(tmp_path):
    fae.write_resume_trace(tmp_path / TR, run_id="20261031T000000Z-0123abcd", state="master-copy",
                           company="Acme", role="Director", model="m", inventory=None, pool=None,
                           res=None, reason="role_map.yaml not found")
    tr = (tmp_path / TR).read_text(encoding="utf-8")
    assert "**state:** master-copy" in tr and "role_map.yaml not found" in tr


def test_new_run_id_format():
    rid = fae.new_run_id()
    assert fae._RUN_ID_RE.search("run_id=" + rid).group(1) == rid


# ── §5.1 publication ──

def make_pair(folder, run_id, master, label):
    """A prepared .new pair; label distinguishes old from new content."""
    folder.mkdir(exist_ok=True)
    r = fae.master_copy_resume(master, folder / (RES + ".new"), run_id)
    t = folder / (TR + ".new")
    t.write_text(f"# trace\n\n**run_id:** `{run_id}`\n\n{label}\n", encoding="utf-8")
    return r, t


def publish(folder, run_id, master, label):
    r, t = make_pair(folder, run_id, master, label)
    return fae.publish_resume_pair(folder, RES, TR, r, t, run_id)


OLD, NEW = "20261030T000000Z-00000001", "20261031T000000Z-00000002"


@pytest.fixture
def setup(tmp_path):
    master = make_master(tmp_path / "m.docx")
    folder = tmp_path / "Acme_Director"
    return master, folder


def state(folder):
    return {p.name for p in folder.iterdir()}, fae.resume_run_id(folder / RES), fae.trace_run_id(folder / TR)


def test_first_run_publishes_cleanly(setup):
    master, folder = setup
    assert publish(folder, NEW, master, "new")["outcome"] == "published"
    names, rr, tr = state(folder)
    assert names == {RES, TR} and rr == tr == NEW


def test_rerun_replaces_pair_and_cleans_up(setup):
    master, folder = setup
    publish(folder, OLD, master, "old")
    assert publish(folder, NEW, master, "new")["outcome"] == "published"
    names, rr, tr = state(folder)
    assert names == {RES, TR} and rr == tr == NEW


def test_failure_after_first_replacement_restores_previous_pair(setup, monkeypatch):
    master, folder = setup
    publish(folder, OLD, master, "old")
    r, t = make_pair(folder, NEW, master, "new")
    real, calls = fae._replace, []

    def flaky(src, dst):
        calls.append((src, dst))
        if str(src).endswith(TR + ".new"):
            raise OSError("disk full")
        real(src, dst)

    monkeypatch.setattr(fae, "_replace", flaky)
    out = fae.publish_resume_pair(folder, RES, TR, r, t, NEW)
    monkeypatch.setattr(fae, "_replace", real)
    assert out["outcome"] == "publication-failed" and "previous resume kept" in out["detail"]
    names, rr, tr = state(folder)
    assert names == {RES, TR} and rr == tr == OLD


def test_unwritable_destination_changes_nothing(setup, monkeypatch):
    master, folder = setup
    publish(folder, OLD, master, "old")
    r, t = make_pair(folder, NEW, master, "new")

    def always_fail(src, dst):
        raise OSError("access denied")

    monkeypatch.setattr(fae, "_replace", always_fail)
    out = fae.publish_resume_pair(folder, RES, TR, r, t, NEW)
    assert out["outcome"] == "publication-failed"
    assert "published" not in out["detail"]
    monkeypatch.undo()
    _, rr, tr = state(folder)
    assert rr == tr == OLD


def test_journal_write_failure_replaces_nothing(setup, monkeypatch):
    master, folder = setup
    publish(folder, OLD, master, "old")
    r, t = make_pair(folder, NEW, master, "new")
    real_write = type(folder).write_text

    def deny(self, *a, **k):
        if self.name == fae.PUBLISH_JOURNAL:
            raise OSError("read-only")
        return real_write(self, *a, **k)

    monkeypatch.setattr(type(folder), "write_text", deny)
    out = fae.publish_resume_pair(folder, RES, TR, r, t, NEW)
    monkeypatch.undo()
    assert out["outcome"] == "publication-failed" and "journal not written" in out["detail"]
    names, rr, tr = state(folder)
    assert names == {RES, TR} and rr == tr == OLD


def test_restore_failure_reports_manual_recovery_and_keeps_journal(setup, monkeypatch):
    master, folder = setup
    publish(folder, OLD, master, "old")
    r, t = make_pair(folder, NEW, master, "new")
    real = fae._replace

    def broken(src, dst):
        if str(src).endswith(".prev") or str(src).endswith(TR + ".new"):
            raise OSError("locked")
        real(src, dst)

    monkeypatch.setattr(fae, "_replace", broken)
    out = fae.publish_resume_pair(folder, RES, TR, r, t, NEW)
    assert out["outcome"] == "publication-failed" and "manual recovery needed" in out["detail"]
    assert (folder / fae.PUBLISH_JOURNAL).exists()


def test_invalid_prepared_pair_is_not_published(setup):
    master, folder = setup
    publish(folder, OLD, master, "old")
    r, t = make_pair(folder, NEW, master, "new")
    t.write_text("# trace without a run id\n", encoding="utf-8")
    out = fae.publish_resume_pair(folder, RES, TR, r, t, NEW)
    assert out["outcome"] == "publication-failed"
    names, rr, tr = state(folder)
    assert names == {RES, TR} and rr == tr == OLD


def interrupt_between(folder, master):
    """Disk state of a crash after the resume was replaced but before the trace."""
    publish(folder, OLD, master, "old")
    r, t = make_pair(folder, NEW, master, "new")
    (folder / fae.PUBLISH_JOURNAL).write_text(json.dumps(
        {"run_id": NEW, "phase": "replacing", "final": [RES, TR], "new": [r.name, t.name]}), encoding="utf-8")
    fae.os.replace(folder / RES, folder / (RES + ".prev"))
    fae.os.replace(folder / TR, folder / (TR + ".prev"))
    fae.os.replace(r, folder / RES)


def test_interruption_between_replacements_rolls_back(setup):
    master, folder = setup
    interrupt_between(folder, master)
    assert fae.resume_pair_problem(folder, RES, TR) == "interrupted resume publication"
    assert fae.recover_resume_publication(folder) == "Previous resume restored after an interrupted publication."
    names, rr, tr = state(folder)
    assert names == {RES, TR} and rr == tr == OLD


def test_interruption_after_both_replacements_rolls_forward(setup):
    master, folder = setup
    interrupt_between(folder, master)
    fae.os.replace(folder / (TR + ".new"), folder / TR)
    assert fae.recover_resume_publication(folder) is None
    names, rr, tr = state(folder)
    assert names == {RES, TR} and rr == tr == NEW


def test_no_journal_means_nothing_to_recover(setup):
    master, folder = setup
    publish(folder, OLD, master, "old")
    assert fae.recover_resume_publication(folder) is None


def test_unreadable_journal_stops_with_instructions(setup):
    master, folder = setup
    publish(folder, OLD, master, "old")
    (folder / fae.PUBLISH_JOURNAL).write_text("{not json", encoding="utf-8")
    with pytest.raises(RuntimeError, match="cannot be read"):
        fae.recover_resume_publication(folder)


# ── startup check helper ──

def test_pair_problem_detects_mismatched_run_ids(setup):
    master, folder = setup
    publish(folder, OLD, master, "old")
    (folder / TR).write_text(f"**run_id:** `{NEW}`\n", encoding="utf-8")
    assert fae.resume_pair_problem(folder, RES, TR) == "resume and trace from different runs"


def test_pair_problem_ignores_pre_adr015_folders(setup):
    master, folder = setup
    folder.mkdir()
    fae.master_copy_resume(master, folder / RES, OLD)  # resume without a trace: an old master copy
    assert fae.resume_pair_problem(folder, RES, TR) is None


def test_pair_problem_ok_for_matching_pair(setup):
    master, folder = setup
    publish(folder, NEW, master, "new")
    assert fae.resume_pair_problem(folder, RES, TR) is None
